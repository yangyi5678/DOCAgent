from __future__ import annotations

import csv
import hashlib
import json
import mimetypes
import re
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Any

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".xlsx", ".csv", ".md", ".txt", ".html", ".htm"}
MAX_FILE_SIZE_BYTES = 200 * 1024 * 1024


@dataclass
class Location:
    page: int | None = None
    sheet: str | None = None
    row_start: int | None = None
    row_end: int | None = None
    bbox: list[float] | None = None


@dataclass
class DocumentMeta:
    file_path: str
    file_name: str
    file_type: str
    mime_type: str
    size_bytes: int
    sha256: str
    page_count: int | None = None
    sheet_count: int | None = None


@dataclass
class Block:
    id: str
    type: str
    text: str
    level: int | None
    location: Location
    order: int
    confidence: float


@dataclass
class Table:
    id: str
    rows: list[list[str]]
    markdown: str
    location: Location
    caption: str | None
    confidence: float


@dataclass
class ImageAsset:
    id: str
    path: str
    location: Location
    caption: str | None
    alt_text: str | None
    confidence: float


@dataclass
class Formula:
    id: str
    latex: str | None
    raw_text: str
    location: Location
    confidence: float


@dataclass
class ParsedDocument:
    meta: DocumentMeta
    blocks: list[Block] = field(default_factory=list)
    tables: list[Table] = field(default_factory=list)
    images: list[ImageAsset] = field(default_factory=list)
    formulas: list[Formula] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


@dataclass
class DocumentChunk:
    id: str
    text: str
    block_ids: list[str]
    page_start: int | None
    page_end: int | None
    sheet: str | None
    token_count: int
    confidence: float


@dataclass
class QualityReport:
    overall: float
    text_quality: float
    table_quality: float
    image_quality: float
    formula_quality: float
    warnings: list[str]


def _safe_resolve(path: str) -> Path:
    return Path(path).expanduser().resolve()


def _safe_under_base(path: Path, base_dir: str | None) -> None:
    if not base_dir:
        return
    base = _safe_resolve(base_dir)
    if base not in [path, *path.parents]:
        raise ValueError("path_traversal_rejected")


def _safe_artifact_dir(export_dir: str, base_dir: str | None = None) -> Path:
    out = _safe_resolve(export_dir)
    _safe_under_base(out, base_dir)
    out.mkdir(parents=True, exist_ok=True)
    return out


def _sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def _detect_encoding(path: Path) -> str:
    try:
        import charset_normalizer

        enc = charset_normalizer.from_bytes(path.read_bytes()[:8192]).best()
        if enc and enc.encoding:
            return enc.encoding
    except Exception:
        pass
    return "utf-8"


def _sanitize_name(name: str) -> str:
    return re.sub(r"[^a-zA-Z0-9._-]", "_", name)


def _table_to_markdown(rows: list[list[str]], limit: int = 20) -> str:
    if not rows:
        return ""
    sample = rows[:limit]
    return "\n".join("|" + "|".join(str(c) for c in r) + "|" for r in sample)


def core_identify_file(file_path: str, base_dir: str | None = None) -> dict[str, Any]:
    p = _safe_resolve(file_path)
    _safe_under_base(p, base_dir)
    if not p.exists():
        raise FileNotFoundError("file_not_found")
    if p.is_dir():
        raise ValueError("directory_not_supported")
    ext = p.suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise ValueError("unsupported_file_type")
    size = p.stat().st_size
    if size > MAX_FILE_SIZE_BYTES:
        raise ValueError("file_too_large")
    mime, _ = mimetypes.guess_type(str(p))
    meta = DocumentMeta(
        file_path=str(p),
        file_name=p.name,
        file_type=ext.lstrip("."),
        mime_type=mime or "application/octet-stream",
        size_bytes=size,
        sha256=_sha256(p),
    )
    if ext == ".pdf":
        try:
            import fitz

            with fitz.open(str(p)) as d:
                meta.page_count = len(d)
        except Exception:
            pass
    if ext == ".xlsx":
        try:
            from openpyxl import load_workbook

            wb = load_workbook(str(p), read_only=True, data_only=False)
            meta.sheet_count = len(wb.sheetnames)
        except Exception:
            pass
    return asdict(meta)


def core_extract_text(file_path: str, base_dir: str | None = None) -> dict[str, Any]:
    info = core_identify_file(file_path, base_dir)
    p = Path(info["file_path"])
    ext = p.suffix.lower()
    blocks: list[Block] = []
    warnings: list[str] = []
    if ext in {".txt", ".md", ".csv", ".html", ".htm"}:
        txt = p.read_text(encoding=_detect_encoding(p), errors="replace")
        if ext in {".html", ".htm"}:
            try:
                from bs4 import BeautifulSoup

                soup = BeautifulSoup(txt, "html.parser")
                main = soup.find("main") or soup.body or soup
                lines = [x.get_text(" ", strip=True) for x in main.find_all(["h1", "h2", "h3", "p", "li", "code"]) if x.get_text(strip=True)]
            except Exception:
                lines = [ln for ln in txt.splitlines() if ln.strip()]
                warnings.append("html_parser_fallback")
        else:
            lines = [ln for ln in txt.splitlines() if ln.strip()]
        order = 1
        for ln in lines:
            btype, level = "paragraph", None
            if ext == ".md" and ln.lstrip().startswith("#"):
                hashes = len(ln) - len(ln.lstrip("#"))
                btype, level = "heading", hashes
                ln = ln[hashes:].strip()
            elif re.match(r"^\s*[-*+]\s+", ln):
                btype = "list"
            blocks.append(Block(f"b{order}", btype, ln, level, Location(), order, 0.95))
            order += 1
    elif ext == ".pdf":
        try:
            import fitz

            with fitz.open(str(p)) as d:
                order = 1
                for pg_no, page in enumerate(d, 1):
                    pdata = page.get_text("dict")
                    page_blocks = []
                    for blk in pdata.get("blocks", []):
                        if blk.get("type") != 0:
                            continue
                        spans = [sp.get("text", "") for ln in blk.get("lines", []) for sp in ln.get("spans", [])]
                        text = " ".join(x for x in spans if x).strip()
                        if not text:
                            continue
                        x0, y0, x1, y1 = blk.get("bbox", [0, 0, 0, 0])
                        col = 0 if x0 < (page.rect.width / 2) else 1
                        page_blocks.append((y0, col, x0, y1, x1, text, blk.get("bbox")))
                    page_blocks.sort(key=lambda x: (x[0], x[1], x[2]))
                    for _, _, _, _, _, text, bbox in page_blocks:
                        blocks.append(Block(f"b{order}", "paragraph", text, None, Location(page=pg_no, bbox=list(bbox)), order, 0.9))
                        order += 1
                    if not page_blocks:
                        warnings.append("pdf_scanned_pages_detected")
                        warnings.append("ocr_fallback_not_available")
        except Exception:
            warnings.append("pdf_text_extraction_failed")
    elif ext == ".docx":
        try:
            from docx import Document

            d = Document(str(p))
            order = 1
            for para in d.paragraphs:
                text = para.text.strip()
                if not text:
                    continue
                style = (para.style.name or "").lower()
                btype = "heading" if "heading" in style else "paragraph"
                level = int(re.sub(r"\D", "", style) or "1") if btype == "heading" else None
                blocks.append(Block(f"b{order}", btype, text, level, Location(), order, 0.95))
                order += 1
        except Exception:
            warnings.append("docx_text_extraction_failed")
    elif ext == ".xlsx":
        try:
            from openpyxl import load_workbook

            wb = load_workbook(str(p), read_only=True, data_only=False)
            order = 1
            for ws in wb.worksheets:
                for ridx, row in enumerate(ws.iter_rows(values_only=True), 1):
                    vals = [str(v) for v in row if v not in (None, "")]
                    if not vals:
                        continue
                    blocks.append(Block(f"b{order}", "paragraph", " | ".join(vals), None, Location(sheet=ws.title, row_start=ridx, row_end=ridx), order, 0.9))
                    order += 1
        except Exception:
            warnings.append("xlsx_text_extraction_failed")
    return {"meta": info, "blocks": [asdict(b) for b in blocks], "warnings": warnings}


def core_extract_tables(file_path: str, base_dir: str | None = None) -> dict[str, Any]:
    info = core_identify_file(file_path, base_dir)
    p = Path(info["file_path"])
    ext = p.suffix.lower()
    tables: list[Table] = []
    warnings: list[str] = []
    if ext == ".csv":
        with p.open("r", encoding=_detect_encoding(p), errors="replace", newline="") as f:
            sample = f.read(8192)
            f.seek(0)
            try:
                dialect = csv.Sniffer().sniff(sample)
            except Exception:
                dialect = csv.excel
            rows = [list(map(str, r)) for r in csv.reader(f, dialect)]
        tables.append(Table("t1", rows, _table_to_markdown(rows), Location(), None, 0.95))
    elif ext == ".xlsx":
        try:
            from openpyxl import load_workbook

            wb = load_workbook(str(p), read_only=True, data_only=False)
            idx = 1
            for ws in wb.worksheets:
                rows = [["" if v is None else str(v) for v in r] for r in ws.iter_rows(values_only=True)]
                tables.append(Table(f"t{idx}", rows, _table_to_markdown(rows), Location(sheet=ws.title), ws.title, 0.9))
                idx += 1
                if ws.merged_cells.ranges:
                    warnings.append("excel_merged_cells_detected")
        except Exception:
            warnings.append("xlsx_table_extraction_failed")
    elif ext == ".md":
        lines = p.read_text(encoding=_detect_encoding(p), errors="replace").splitlines()
        raw = [ln for ln in lines if "|" in ln]
        if raw:
            rows = [[c.strip() for c in ln.strip("|").split("|")] for ln in raw]
            tables.append(Table("t1", rows, _table_to_markdown(rows), Location(), None, 0.7))
    elif ext in {".html", ".htm"}:
        try:
            from bs4 import BeautifulSoup

            soup = BeautifulSoup(p.read_text(encoding=_detect_encoding(p), errors="replace"), "html.parser")
            idx = 1
            for t in soup.find_all("table"):
                rows = []
                for tr in t.find_all("tr"):
                    rows.append([c.get_text(" ", strip=True) for c in tr.find_all(["th", "td"])])
                tables.append(Table(f"t{idx}", rows, _table_to_markdown(rows), Location(), None, 0.85))
                idx += 1
        except Exception:
            warnings.append("html_table_extraction_failed")
    elif ext == ".pdf":
        warnings.append("pdf_table_extraction_not_enabled")
    elif ext == ".docx":
        try:
            from docx import Document

            d = Document(str(p))
            for idx, t in enumerate(d.tables, 1):
                rows = [[c.text for c in r.cells] for r in t.rows]
                tables.append(Table(f"t{idx}", rows, _table_to_markdown(rows), Location(), None, 0.85))
        except Exception:
            warnings.append("docx_table_extraction_failed")
    return {"tables": [asdict(t) for t in tables], "warnings": warnings}


def core_extract_images(file_path: str, artifact_dir: str | None = None, base_dir: str | None = None) -> dict[str, Any]:
    info = core_identify_file(file_path, base_dir)
    p = Path(info["file_path"])
    out = _safe_artifact_dir(artifact_dir, base_dir) if artifact_dir else None
    images: list[ImageAsset] = []
    warnings: list[str] = []
    if p.suffix.lower() == ".pdf":
        try:
            import fitz

            with fitz.open(str(p)) as d:
                idx = 1
                for pn, page in enumerate(d, 1):
                    for img in page.get_images(full=True):
                        if not out:
                            warnings.append("artifact_dir_required_for_image_export")
                            break
                        xref = img[0]
                        pix = fitz.Pixmap(d, xref)
                        name = f"pdf_p{pn}_img{idx}.png"
                        fp = out / _sanitize_name(name)
                        pix.save(str(fp))
                        images.append(ImageAsset(f"img{idx}", str(fp), Location(page=pn), None, None, 0.85))
                        idx += 1
        except Exception:
            warnings.append("pdf_image_extraction_failed")
    elif p.suffix.lower() == ".docx":
        warnings.append("docx_image_extraction_not_enabled")
    return {"images": [asdict(i) for i in images], "warnings": warnings}


def core_extract_formulas(file_path: str, base_dir: str | None = None) -> dict[str, Any]:
    info = core_identify_file(file_path, base_dir)
    p = Path(info["file_path"])
    warnings = []
    formulas: list[Formula] = []
    if p.suffix.lower() in {".pdf", ".docx"}:
        warnings.append("formula_low_confidence")
        warnings.append("formula_extraction_partial")
    return {"formulas": [asdict(f) for f in formulas], "warnings": warnings}


def core_chunk(parsed_document: dict[str, Any], chunk_size: int = 800, overlap: int = 100) -> dict[str, Any]:
    blocks = parsed_document.get("blocks", [])
    out: list[DocumentChunk] = []
    current_text: list[str] = []
    current_ids: list[str] = []
    idx = 1
    start_page = end_page = None
    start_sheet = None
    for b in blocks:
        btxt = b.get("text", "")
        btype = b.get("type", "paragraph")
        if not btxt:
            continue
        btoks = len(btxt.split())
        ctoks = len(" ".join(current_text).split())
        hard_boundary = btype in {"table_ref", "caption", "formula_ref"}
        if current_text and (ctoks + btoks > chunk_size or hard_boundary):
            txt = "\n".join(current_text)
            out.append(DocumentChunk(f"c{idx}", txt, current_ids.copy(), start_page, end_page or start_page, start_sheet, len(txt.split()), 0.9))
            idx += 1
            if overlap > 0 and current_text:
                current_text = current_text[-1:]
                current_ids = current_ids[-1:]
            else:
                current_text, current_ids = [], []
                start_page = end_page = None
                start_sheet = None
        if not current_text:
            start_page = b.get("location", {}).get("page")
            start_sheet = b.get("location", {}).get("sheet")
        end_page = b.get("location", {}).get("page", end_page)
        current_text.append(btxt)
        current_ids.append(b.get("id", ""))
    if current_text:
        txt = "\n".join(current_text)
        out.append(DocumentChunk(f"c{idx}", txt, current_ids, start_page, end_page or start_page, start_sheet, len(txt.split()), 0.9))
    return {"chunks": [asdict(c) for c in out]}


def core_evaluate_quality(parsed_document: dict[str, Any]) -> dict[str, Any]:
    blocks = parsed_document.get("blocks", [])
    tables = parsed_document.get("tables", [])
    images = parsed_document.get("images", [])
    formulas = parsed_document.get("formulas", [])
    warnings = list(parsed_document.get("warnings", []))

    if not blocks:
        warnings.append("empty_content_ratio_high")
    if any("�" in (b.get("text") or "") for b in blocks):
        warnings.append("garbled_text_detected")
    if any(t.get("confidence", 1.0) < 0.6 for t in tables):
        warnings.append("table_low_confidence")
    if any(f.get("confidence", 1.0) < 0.6 for f in formulas):
        warnings.append("formula_low_confidence")

    text_quality = 0.9 if blocks else 0.2
    table_quality = (sum(t.get("confidence", 0.0) for t in tables) / len(tables)) if tables else 0.8
    image_quality = (sum(i.get("confidence", 0.0) for i in images) / len(images)) if images else 0.8
    formula_quality = (sum(f.get("confidence", 0.0) for f in formulas) / len(formulas)) if formulas else 0.6
    overall = round((text_quality + table_quality + image_quality + formula_quality) / 4.0, 3)
    return {"quality": asdict(QualityReport(overall, text_quality, table_quality, image_quality, formula_quality, sorted(set(warnings))))}


def core_export_artifacts(parsed_document: dict[str, Any], export_dir: str, base_dir: str | None = None) -> dict[str, Any]:
    out = _safe_artifact_dir(export_dir, base_dir)
    parsed_json = out / "parsed_document.json"
    parsed_json.write_text(json.dumps(parsed_document, ensure_ascii=False, indent=2), encoding="utf-8")

    blocks_md = out / "document_markdown.md"
    blocks_md.write_text("\n\n".join(b.get("text", "") for b in parsed_document.get("blocks", [])), encoding="utf-8")

    chunks = core_chunk(parsed_document).get("chunks", [])
    chunks_jsonl = out / "chunks.jsonl"
    with chunks_jsonl.open("w", encoding="utf-8") as f:
        for c in chunks:
            f.write(json.dumps(c, ensure_ascii=False) + "\n")

    tables_csv = []
    for idx, t in enumerate(parsed_document.get("tables", []), 1):
        p = out / f"table_{idx}.csv"
        with p.open("w", encoding="utf-8", newline="") as f:
            writer = csv.writer(f)
            writer.writerows(t.get("rows", []))
        tables_csv.append(str(p))

    quality = core_evaluate_quality(parsed_document)["quality"]
    quality_report = out / "quality_report.json"
    quality_report.write_text(json.dumps(quality, ensure_ascii=False, indent=2), encoding="utf-8")

    return {
        "artifacts": {
            "parsed_json": str(parsed_json),
            "document_markdown": str(blocks_md),
            "chunks_jsonl": str(chunks_jsonl),
            "tables_csv": tables_csv,
            "images_dir": str(out),
            "quality_report": str(quality_report),
        }
    }


def core_parse(file_path: str, export_dir: str | None = None, base_dir: str | None = None) -> dict[str, Any]:
    text = core_extract_text(file_path, base_dir)
    tables = core_extract_tables(file_path, base_dir)
    images = core_extract_images(file_path, export_dir, base_dir) if export_dir else {"images": [], "warnings": ["image_extraction_skipped_no_artifact_dir"]}
    formulas = core_extract_formulas(file_path, base_dir)

    parsed_dict = {
        "meta": text["meta"],
        "blocks": text.get("blocks", []),
        "tables": tables.get("tables", []),
        "images": images.get("images", []),
        "formulas": formulas.get("formulas", []),
        "warnings": text.get("warnings", []) + tables.get("warnings", []) + images.get("warnings", []) + formulas.get("warnings", []),
    }
    chunks = core_chunk(parsed_dict)["chunks"]
    quality = core_evaluate_quality(parsed_dict)["quality"]
    artifacts = core_export_artifacts(parsed_dict, export_dir, base_dir)["artifacts"] if export_dir else None
    return {"parsed_document": parsed_dict, "chunks": chunks, "quality": quality, "artifacts": artifacts}
