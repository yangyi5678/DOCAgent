from pathlib import Path

import pytest

from tools.doc_tools import (
    chunk_tool,
    evaluate_quality_tool,
    export_artifacts_tool,
    identify_file_tool,
    parse_tool,
)


def test_identify_file_missing_unsupported_directory(tmp_path: Path):
    p = tmp_path / "a.txt"
    p.write_text("hello")
    ok = identify_file_tool(str(p), base_dir=str(tmp_path))
    assert ok["success"]

    missing = identify_file_tool(str(tmp_path / "missing.txt"), base_dir=str(tmp_path))
    assert not missing["success"] and "file_not_found" in missing["error"]

    unsupported = tmp_path / "bad.bin"
    unsupported.write_bytes(b"x")
    bad = identify_file_tool(str(unsupported), base_dir=str(tmp_path))
    assert not bad["success"] and "unsupported_file_type" in bad["error"]

    dres = identify_file_tool(str(tmp_path), base_dir=str(tmp_path))
    assert not dres["success"] and "directory_not_supported" in dres["error"]


def test_parse_txt_md_csv_html(tmp_path: Path):
    (tmp_path / "a.txt").write_text("Title\n\nBody")
    (tmp_path / "a.md").write_text("# H1\n- item")
    (tmp_path / "a.csv").write_text("c1,c2\n1,2")
    (tmp_path / "a.html").write_text("<html><body><main><h1>T</h1><p>P</p></main></body></html>")

    for n in ["a.txt", "a.md", "a.csv", "a.html"]:
        res = parse_tool(str(tmp_path / n), base_dir=str(tmp_path))
        assert res["success"]
        assert "parsed_document" in res["data"]


def test_parse_xlsx_docx_pdf_optional(tmp_path: Path):
    try:
        from openpyxl import Workbook
    except Exception:
        pytest.skip("openpyxl unavailable")
    wb = Workbook()
    ws = wb.active
    ws.append(["h1", "h2"])
    ws.append([1, 2])
    xlsx = tmp_path / "a.xlsx"
    wb.save(xlsx)
    assert parse_tool(str(xlsx), base_dir=str(tmp_path))["success"]

    try:
        from docx import Document
    except Exception:
        pytest.skip("python-docx unavailable")
    d = Document()
    d.add_heading("Head", level=1)
    d.add_paragraph("Body")
    docx = tmp_path / "a.docx"
    d.save(docx)
    assert parse_tool(str(docx), base_dir=str(tmp_path))["success"]

    try:
        import fitz
    except Exception:
        pytest.skip("pymupdf unavailable")
    pdf = tmp_path / "a.pdf"
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Hello PDF")
    doc.save(pdf)
    doc.close()
    assert parse_tool(str(pdf), base_dir=str(tmp_path))["success"]


def test_chunk_quality_export_path_traversal(tmp_path: Path):
    parsed_doc = {
        "meta": {},
        "blocks": [
            {"id": "b1", "type": "heading", "text": "H", "location": {"page": 1}},
            {"id": "b2", "type": "paragraph", "text": "x y z", "location": {"page": 1}},
        ],
        "tables": [],
        "images": [],
        "formulas": [],
        "warnings": [],
    }
    ch = chunk_tool(parsed_doc, chunk_size=2, overlap=0)
    assert ch["success"] and len(ch["data"]["chunks"]) >= 1

    q = evaluate_quality_tool(parsed_doc)
    assert q["success"] and "overall" in q["data"]["quality"]

    out = tmp_path / "out"
    ex = export_artifacts_tool(parsed_doc, str(out), base_dir=str(tmp_path))
    assert ex["success"]
    assert (out / "parsed_document.json").exists()
    assert (out / "document_markdown.md").exists()
    assert (out / "chunks.jsonl").exists()
    assert (out / "quality_report.json").exists()

    evil = identify_file_tool(str(tmp_path / "../x.txt"), base_dir=str(tmp_path))
    assert not evil["success"] and "path_traversal_rejected" in evil["error"]
